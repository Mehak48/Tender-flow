"""Stage 9: exports are real files, re-opened and checked; the uploaded template never changes."""

from __future__ import annotations

import io
import re

import pytest
from docx import Document
from openpyxl import load_workbook

from app.config import load_settings
from app.export.pdf_export import find_soffice
from tests.conftest import ready_project


def _docx_text(data: bytes) -> str:
    doc = Document(io.BytesIO(data))
    parts = [p.text for p in doc.paragraphs]
    for t in doc.tables:
        for row in t.rows:
            parts.extend(c.text for c in row.cells)
    for s in doc.sections:
        parts.extend(p.text for p in s.header.paragraphs)
        parts.extend(p.text for p in s.footer.paragraphs)
    return "\n".join(parts)


def _make_ready(client, view: dict) -> dict:  # type: ignore[no-untyped-def]
    pid = view["id"]
    rack = next(i for i in view["quotation"]["line_items"] if i["unit_price"] is None)
    client.put(
        f"/api/projects/{pid}/quotation",
        json={
            "operations": [
                {"op": "set_item", "line_item_id": rack["id"], "attribute": "unit_price", "value": "95000"},
                {"op": "set_currency", "currency": "PKR"},
                {"op": "set_field", "field_key": "company_name", "value": "Test Supplier Ltd"},
                {"op": "set_field", "field_key": "company_address", "value": "1 Test Street"},
                {"op": "set_field", "field_key": "company_phone", "value": "000"},
                {"op": "set_field", "field_key": "company_email", "value": "sales@test.example"},
                {"op": "set_field", "field_key": "company_registration", "value": "REG-1"},
                {"op": "set_field", "field_key": "prepared_by", "value": "A. Tester"},
            ]
        },
    )
    client.post(f"/api/projects/{pid}/quotation/fix", json={"fix": "confirm_prices"})
    r = client.post(f"/api/projects/{pid}/quotation/fix", json={"fix": "confirm_terms"})
    for issue in r.json()["validation"]["issues"]:
        if issue["severity"] == "warning":
            r = client.post(f"/api/projects/{pid}/validation/acknowledge", json={"issue_id": issue["id"]})
    assert r.json()["validation"]["ready"] is True, r.json()["validation"]
    return r.json()


def test_final_export_is_blocked_until_ready_but_draft_is_allowed(client, docs, settings) -> None:  # type: ignore[no-untyped-def]
    view = ready_project(client, docs)
    pid = view["id"]
    r = client.post(f"/api/projects/{pid}/export", json={"format": "docx"})
    assert r.status_code == 409 and "draft" in r.json()["error"]["message"]
    r = client.post(f"/api/projects/{pid}/export", json={"format": "docx", "draft": True})
    assert r.status_code == 200, r.text
    artifact = r.json()["artifact"]
    assert artifact["draft"] is True and artifact["filename"].endswith("_DRAFT.docx")
    data = client.get(f"/api/projects/{pid}/exports/{artifact['id']}/download").content
    text = _docx_text(data)
    assert "DRAFT" in text
    assert "[MISSING: Unit price]" in text  # the rack has no price and none was invented
    assert "[MISSING: Company name]" in text
    assert not re.search(r"\{\{\s*[a-z_.]+\s*\??\s*\}\}", text)
    names = {c["name"]: c["passed"] for c in artifact["checks"]}
    assert names["Original template file unchanged"] is True
    assert names["Missing values are visibly marked"] is True


def test_docx_export_contains_the_quotation(client, docs) -> None:  # type: ignore[no-untyped-def]
    view = _make_ready(client, ready_project(client, docs))
    pid = view["id"]
    r = client.post(f"/api/projects/{pid}/export", json={"format": "docx"})
    assert r.status_code == 200, r.text
    artifact = r.json()["artifact"]
    assert artifact["draft"] is False
    assert all(c["passed"] for c in artifact["checks"]), artifact["checks"]
    data = client.get(f"/api/projects/{pid}/exports/{artifact['id']}/download").content
    assert len(data) == artifact["size_bytes"]
    doc = Document(io.BytesIO(data))
    text = _docx_text(data)
    for expected in ("NPI/PROC/2026/017", "1,564,000.00", "95,000.00", "Test Supplier Ltd", "12 February 2027"):
        assert expected in text
    assert "DRAFT" not in text
    items_table = doc.tables[1]
    rows = [[c.text for c in row.cells] for row in items_table.rows]
    assert rows[1][1] == "Managed network switch" and rows[1][3] == "4" and rows[1][6] == "740,000.00"
    assert len(rows) == 1 + 5 + 3  # header + items + subtotal/tax/total
    assert "Signature: ____" in text  # signatures are never filled


def test_xlsx_template_export_writes_numbers(client, docs) -> None:  # type: ignore[no-untyped-def]
    view = _make_ready(client, ready_project(client, docs, template="template_xlsx"))
    pid = view["id"]
    assert view["export_formats"] == ["xlsx", "pdf"]
    r = client.post(f"/api/projects/{pid}/export", json={"format": "xlsx"})
    assert r.status_code == 200, r.text
    data = client.get(f"/api/projects/{pid}/exports/{r.json()['artifact']['id']}/download").content
    ws = load_workbook(io.BytesIO(data)).active
    assert ws["B13"].value == "Managed network switch"
    assert float(ws["D13"].value) == 4 and float(ws["G13"].value) == 740000
    assert float(ws["G21"].value) == 1564000  # {{subtotal}} written as a number
    assert ws["C8"].value == "NPI/PROC/2026/017"
    assert ws["B32"].value == "A. Tester"  # "Prepared by" label-blank cell filled
    r = client.post(f"/api/projects/{pid}/export", json={"format": "docx"})
    assert r.status_code == 422


def test_placeholder_row_template_and_csv_prices(client, docs) -> None:  # type: ignore[no-untyped-def]
    view = ready_project(client, docs, template="template_rows_docx", prices="prices_csv", tender="tender_docx")
    pid = view["id"]
    r = client.post(f"/api/projects/{pid}/export", json={"format": "docx", "draft": True})
    data = client.get(f"/api/projects/{pid}/exports/{r.json()['artifact']['id']}/download").content
    rows = [[c.text for c in row.cells] for row in Document(io.BytesIO(data)).tables[0].rows]
    assert rows[1] == ["Managed network switch", "4", "185,000.00", "740,000.00"]
    assert rows[4][2] == "[MISSING: Unit price]"


def test_preview_returns_a_valid_docx(client, docs) -> None:  # type: ignore[no-untyped-def]
    view = ready_project(client, docs)
    r = client.get(f"/api/projects/{view['id']}/preview/docx")
    assert r.status_code == 200 and r.content[:2] == b"PK"
    assert "NPI/PROC/2026/017" in _docx_text(r.content)
    assert client.get(f"/api/projects/{view['id']}").json()["exports"] == []  # previews are not exports


def test_pdf_unavailable_is_reported_honestly(client, docs, monkeypatch) -> None:  # type: ignore[no-untyped-def]
    monkeypatch.setattr("app.export.pdf_export.find_soffice", lambda _settings: None)
    view = ready_project(client, docs)
    r = client.post(f"/api/projects/{view['id']}/export", json={"format": "pdf", "draft": True})
    assert r.status_code == 422 and r.json()["error"]["code"] == "pdf_unavailable"
    assert client.get(f"/api/projects/{view['id']}").json()["exports"] == []


@pytest.mark.skipif(find_soffice(load_settings()) is None, reason="LibreOffice (soffice) is not installed")
def test_pdf_export_with_libreoffice(client, docs) -> None:  # type: ignore[no-untyped-def]
    view = _make_ready(client, ready_project(client, docs))
    r = client.post(f"/api/projects/{view['id']}/export", json={"format": "pdf"})
    assert r.status_code == 200, r.text
    artifact = r.json()["artifact"]
    assert all(c["passed"] for c in artifact["checks"]), artifact["checks"]
    data = client.get(f"/api/projects/{view['id']}/exports/{artifact['id']}/download").content
    assert data[:5] == b"%PDF-"


def test_export_download_and_delete(client, docs, settings) -> None:  # type: ignore[no-untyped-def]
    view = ready_project(client, docs)
    pid = view["id"]
    artifact = client.post(f"/api/projects/{pid}/export", json={"format": "docx", "draft": True}).json()["artifact"]
    path = settings.data_dir / "projects" / pid / "exports" / f"{artifact['id']}.docx"
    assert path.is_file()
    r = client.get(f"/api/projects/{pid}/exports/{artifact['id']}/download")
    assert artifact["filename"] in r.headers["content-disposition"]
    assert client.delete(f"/api/projects/{pid}/exports/{artifact['id']}").status_code == 200
    assert not path.exists()
    assert client.get(f"/api/projects/{pid}/exports/{artifact['id']}/download").status_code == 404
