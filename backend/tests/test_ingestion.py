"""Stage 1: upload validation and safe storage."""

from __future__ import annotations

import io
import re
import time
import zipfile
from pathlib import Path

import pytest

from app.errors import AppError
from app.ingestion import sanitize_filename, validate_upload
from tests.conftest import new_project, upload

MAX = 2 * 1024 * 1024


@pytest.mark.parametrize(
    ("raw", "expected"),
    [
        ("../../etc/passwd.pdf", "passwd.pdf"),
        ("C:\\Users\\me\\Desktop\\tender.pdf", "tender.pdf"),
        ("ten\x00der\x07.pdf", "tender.pdf"),
        ("", "document"),
        ("   ", "document"),
        ("<script>alert(1)</script>.docx", "script_.docx"),  # "/" is a path separator: only the basename is kept
    ],
)
def test_sanitize_filename(raw: str, expected: str) -> None:
    assert sanitize_filename(raw) == expected


def test_sanitize_filename_bounds_length() -> None:
    name = sanitize_filename("a" * 500 + ".docx")
    assert len(name) <= 120 and name.endswith(".docx")


def test_rejects_unsupported_extension() -> None:
    with pytest.raises(AppError) as exc:
        validate_upload("tender", "tender.exe", b"MZ....", MAX)
    assert exc.value.code == "unsupported_format" and exc.value.status_code == 415


def test_legacy_doc_gets_helpful_message() -> None:
    with pytest.raises(AppError) as exc:
        validate_upload("tender", "tender.doc", b"\xd0\xcf\x11\xe0", MAX)
    assert ".docx" in exc.value.message


def test_rejects_wrong_role_format_combination() -> None:
    with pytest.raises(AppError):
        validate_upload("template", "template.pdf", b"%PDF-1.4", MAX)


def test_rejects_empty_and_oversized() -> None:
    with pytest.raises(AppError) as empty:
        validate_upload("tender", "t.pdf", b"", MAX)
    assert empty.value.code == "empty_file"
    with pytest.raises(AppError) as big:
        validate_upload("tender", "t.pdf", b"%PDF-" + b"0" * MAX, MAX)
    assert big.value.status_code == 413


def test_rejects_fake_pdf_and_fake_docx() -> None:
    with pytest.raises(AppError) as pdf:
        validate_upload("tender", "t.pdf", b"this is not a pdf", MAX)
    assert pdf.value.code == "malformed_file"
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as zf:
        zf.writestr("hello.txt", "not a word document")
    with pytest.raises(AppError) as docx:
        validate_upload("tender", "t.docx", buf.getvalue(), MAX)
    assert docx.value.code == "malformed_file"


def test_rejects_zip_bomb_like_container() -> None:
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", compression=zipfile.ZIP_DEFLATED) as zf:
        zf.writestr("[Content_Types].xml", "<x/>")
        zf.writestr("word/document.xml", "<x/>")
        zf.writestr("big.bin", b"\0" * (310 * 1024 * 1024))
    with pytest.raises(AppError) as exc:
        validate_upload("tender", "t.docx", buf.getvalue(), 50 * 1024 * 1024)
    assert exc.value.code == "unsafe_file"


def test_api_upload_errors_are_clear(client, docs) -> None:  # type: ignore[no-untyped-def]
    pid = new_project(client)
    r = client.post(
        f"/api/projects/{pid}/documents",
        data={"role": "tender"},
        files={"file": ("x.png", b"\x89PNG....", "image/png")},
    )
    assert r.status_code == 415
    assert "OCR" in r.json()["error"]["message"]
    r = client.post(
        f"/api/projects/{pid}/documents",
        data={"role": "tender"},
        files={"file": ("huge.pdf", b"%PDF-" + b"0" * (3 * 1024 * 1024), "application/pdf")},
    )
    assert r.status_code == 413
    r = client.post(
        f"/api/projects/{pid}/documents", data={"role": "nonsense"}, files={"file": ("a.pdf", b"%PDF-1", "x")}
    )
    assert r.status_code == 400
    # A malformed PDF that passes the header check is rejected at parse time and not kept on disk.
    r = client.post(
        f"/api/projects/{pid}/documents",
        data={"role": "tender"},
        files={"file": ("broken.pdf", b"%PDF-1.7\n garbage garbage", "application/pdf")},
    )
    assert r.status_code in (400, 422)
    assert client.get(f"/api/projects/{pid}").json()["documents"] == []


def test_uploads_are_stored_under_generated_ids(client, docs, settings) -> None:  # type: ignore[no-untyped-def]
    pid = new_project(client)
    r = upload(client, pid, "tender", docs["tender_pdf"], filename="../../../evil name.pdf")
    assert r.status_code == 201
    record = r.json()["documents"][0]
    assert record["filename"] == "evil name.pdf"
    stored = list((settings.data_dir / "projects" / pid / "uploads").iterdir())
    assert [p.name for p in stored] == [f"{record['id']}.pdf"]
    assert str(settings.data_dir) not in r.text  # no filesystem paths in API responses


def test_replacing_a_document_removes_the_old_file(client, docs, settings) -> None:  # type: ignore[no-untyped-def]
    pid = new_project(client)
    first = upload(client, pid, "tender", docs["tender_pdf"]).json()["documents"][0]["id"]
    second = upload(client, pid, "tender", docs["tender_docx"]).json()["documents"]
    assert len(second) == 1 and second[0]["id"] != first
    names = [p.name for p in (settings.data_dir / "projects" / pid / "uploads").iterdir()]
    assert names == [f"{second[0]['id']}.docx"]


_BOMB_DTD = '<!ENTITY a0 "lol">' + "".join(f'<!ENTITY a{i} "{f"&a{i - 1};" * 10}">' for i in range(1, 10))


def _with_entity_bomb(src: Path, needle: str, out: Path) -> Path:
    """Copy an OOXML file and turn the part containing ``needle`` into a "billion laughs" XML bomb."""
    injected = False
    with zipfile.ZipFile(src) as zin, zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as zout:
        for item in zin.infolist():
            data = zin.read(item.filename)
            text = data.decode("utf-8", errors="ignore")
            if not injected and needle in text and item.filename.endswith(".xml"):
                head, sep, rest = text.partition("?>")
                head, rest = (head + sep, rest) if sep else ("", text)
                root = re.search(r"<([A-Za-z_][\w:.-]*)[\s>]", rest)
                assert root
                text = head + f"<!DOCTYPE {root.group(1)} [{_BOMB_DTD}]>" + rest.replace(needle, "&a9;", 1)
                data, injected = text.encode("utf-8"), True
            zout.writestr(item, data)
    assert injected
    return out


def test_xml_entity_bombs_are_not_expanded(client, settings, tmp_path) -> None:  # type: ignore[no-untyped-def]
    """Uploaded Office files are untrusted: entity expansion must not hang the server or bloat the text."""
    from docx import Document
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.append(["Item", "Unit price"])
    ws.append(["BOMBCELL", 5])
    wb.save(tmp_path / "prices.xlsx")
    doc = Document()
    doc.add_paragraph("Tender reference: TF-0001")
    doc.add_paragraph("BOMBPARA")
    doc.save(tmp_path / "tender.docx")

    pid = new_project(client)
    accepted: set[str] = set()
    for role, src, needle in (
        ("price_list", tmp_path / "prices.xlsx", "BOMBCELL"),
        ("tender", tmp_path / "tender.docx", "BOMBPARA"),
    ):
        bomb = _with_entity_bomb(src, needle, tmp_path / f"bomb-{src.name}")
        started = time.monotonic()
        r = upload(client, pid, role, bomb)
        assert time.monotonic() - started < 10
        assert r.status_code in (201, 400, 422), r.text
        assert "lollollol" not in r.text
        if r.status_code == 201:
            accepted.add(role)
        else:  # rejected with a clear, user-facing message
            assert r.json()["error"]["code"].startswith("unreadable")
    # Rejected files are not kept, neither in the project nor on disk.
    documents = client.get(f"/api/projects/{pid}").json()["documents"]
    assert {d["role"] for d in documents} == accepted
    stored = list((settings.data_dir / "projects" / pid / "uploads").glob("*"))
    assert len(stored) == len(accepted)


def test_invalid_ids_cannot_reach_the_filesystem(client) -> None:  # type: ignore[no-untyped-def]
    for bad in ("..", "prj_..%2F..%2Fetc", "prj_ZZZZZZZZZZZZ", "prj_123"):
        assert client.get(f"/api/projects/{bad}").status_code == 404
