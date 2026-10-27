"""Export orchestration: generate the file, re-open it with the matching parser and record checks.

An export artifact is only reported when the file exists on disk and could be re-opened.
"""

from __future__ import annotations

import re
import shutil
import tempfile
from decimal import Decimal, InvalidOperation
from pathlib import Path

import pdfplumber

from ..config import Settings
from ..errors import AppError
from ..models import ExportArtifact, ExportCheck, ExportFormat, Project, ValidationReport
from ..storage import ProjectStore, new_id
from .common import Resolver, sha256_file
from .docx_export import docx_text, export_docx
from .pdf_export import convert_to_pdf
from .xlsx_export import export_xlsx, xlsx_text


def _numbers(text: str) -> set[Decimal]:
    found: set[Decimal] = set()
    for m in re.findall(r"\d[\d,]*(?:\.\d+)?", text):
        try:
            found.add(Decimal(m.replace(",", "")).normalize())
        except InvalidOperation:
            continue
    return found


def allowed_formats(project: Project) -> list[ExportFormat]:
    if project.template is None:
        return []
    return ["docx", "pdf"] if project.template.format == "docx" else ["xlsx", "pdf"]


def output_filename(project: Project, fmt: str, draft: bool) -> str:
    q = project.quotation
    number = q.fields.get("quotation_number").value if q and q.fields.get("quotation_number") else None  # type: ignore[union-attr]
    base = re.sub(r"[^A-Za-z0-9._-]+", "_", number or "Quotation").strip("._-") or "Quotation"
    return f"{base[:60]}{'_DRAFT' if draft else ''}.{fmt}"


def _template_record(project: Project):
    return next((d for d in project.documents if d.role == "template"), None)


def generate_file(
    store: ProjectStore, settings: Settings, project: Project, fmt: ExportFormat, draft: bool, out_path: Path
) -> tuple[list[str], Resolver]:
    q, template, mapping = project.quotation, project.template, project.mapping
    record = _template_record(project)
    if q is None or template is None or mapping is None or record is None:
        raise AppError("not_ready", "Generate the quotation (with a template and mapping) before exporting.", 409)
    template_path = store.upload_path(project.id, record.id, record.format)
    if not template_path.is_file():
        raise AppError("template_missing", "The stored template file is missing. Upload the template again.", 409)
    resolver = Resolver(q=q, mapping=mapping, template=template)
    if fmt == "pdf":
        with tempfile.TemporaryDirectory(prefix="tf-export-") as tmp:
            intermediate = Path(tmp) / f"quotation.{template.format}"
            if template.format == "docx":
                warnings = export_docx(template_path, intermediate, resolver, draft)
            else:
                warnings = export_xlsx(template_path, intermediate, resolver, draft)
            convert_to_pdf(intermediate, out_path, settings)
        return warnings, resolver
    if fmt == "docx" and template.format == "docx":
        return export_docx(template_path, out_path, resolver, draft), resolver
    if fmt == "xlsx" and template.format == "xlsx":
        return export_xlsx(template_path, out_path, resolver, draft), resolver
    raise AppError(
        "unsupported_export", f"A {template.format.upper()} template cannot be exported as {fmt.upper()}.", 422
    )


def _verify(project: Project, fmt: str, path: Path, resolver: Resolver) -> list[ExportCheck]:
    checks: list[ExportCheck] = []
    if not path.is_file() or path.stat().st_size == 0:
        raise AppError("export_failed", "The export file was not created.", 500)
    try:
        if fmt == "docx":
            text = docx_text(path)
        elif fmt == "xlsx":
            text = xlsx_text(path)
        else:
            with pdfplumber.open(str(path)) as pdf:
                text = "\n".join((page.extract_text() or "") for page in pdf.pages)
        checks.append(ExportCheck(name=f"File opens as {fmt.upper()}", passed=True))
    except Exception as exc:
        raise AppError("export_failed", f"The generated {fmt.upper()} file could not be re-opened.", 500) from exc
    leftover = re.findall(r"\{\{\s*[A-Za-z][A-Za-z0-9_.]*\s*\??\s*\}\}", text)
    checks.append(
        ExportCheck(
            name="No unfilled {{placeholders}} left",
            passed=not leftover,
            detail=", ".join(sorted(set(leftover))[:5]) if leftover else None,
        )
    )
    q = project.quotation
    assert q is not None
    flat = re.sub(r"\s+", " ", text)
    numbers_in_doc = _numbers(flat)
    for key, label in (("tender_reference", "Tender reference"), ("total", "Quotation total")):
        qf = q.fields.get(key)
        mapped = any(r.source_key == key for r in resolver.mapping.fields)
        if qf and qf.value and mapped:
            value_flat = re.sub(r"\s+", " ", qf.value)
            present = value_flat in flat
            if not present and key == "total" and q.total is not None:
                present = q.total.normalize() in numbers_in_doc
            checks.append(ExportCheck(name=f"{label} appears in the document", passed=present, detail=qf.value))
    compact = re.sub(r"\s+", "", text)
    words = set(re.findall(r"[a-z0-9]+", text.lower()))

    def written(desc: str) -> bool:
        if re.sub(r"\s+", "", desc[:25]) in compact:
            return True
        if fmt == "pdf":  # PDF text of wrapped table cells interleaves columns: compare words instead
            return set(re.findall(r"[a-z0-9]+", desc.lower()[:60])) <= words
        return False

    missing_desc = [it for it in q.line_items if it.description and not written(it.description)]
    if q.line_items and resolver.template.items_table is not None:
        checks.append(
            ExportCheck(
                name=f"All {len(q.line_items)} line items written",
                passed=not missing_desc,
                detail=None if not missing_desc else f"{len(missing_desc)} not found",
            )
        )
    if resolver.missing_written:
        checks.append(
            ExportCheck(
                name="Missing values are visibly marked",
                passed=True,
                detail=f"{len(set(resolver.missing_written))} field(s) marked [MISSING: …]",
            )
        )
    return checks


def run_export(
    store: ProjectStore, settings: Settings, project: Project, fmt: ExportFormat, draft: bool, report: ValidationReport
) -> ExportArtifact:
    if fmt not in allowed_formats(project):
        raise AppError("unsupported_export", f"{fmt.upper()} export is not available for this template.", 422)
    if not draft and not report.ready:
        raise AppError(
            "not_ready",
            f"The quotation is not ready: {report.errors} error(s) and {report.unacknowledged_warnings} "
            "unacknowledged warning(s). Resolve them, or export a draft for internal review.",
            409,
        )
    record = _template_record(project)
    template_path = store.upload_path(project.id, record.id, record.format) if record else None
    template_hash_before = sha256_file(template_path) if template_path and template_path.is_file() else None

    export_id = new_id("exp")
    final_path = store.export_path(project.id, export_id, fmt)
    with tempfile.TemporaryDirectory(prefix="tf-out-") as tmp:
        tmp_path = Path(tmp) / f"out.{fmt}"
        warnings, resolver = generate_file(store, settings, project, fmt, draft, tmp_path)
        checks = _verify(project, fmt, tmp_path, resolver)
        final_path.parent.mkdir(parents=True, exist_ok=True)
        shutil.move(str(tmp_path), final_path)

    if template_path is not None and template_hash_before is not None:
        unchanged = sha256_file(template_path) == template_hash_before == (record.sha256 if record else None)
        checks.append(ExportCheck(name="Original template file unchanged", passed=unchanged))
    return ExportArtifact(
        id=export_id,
        format=fmt,
        filename=output_filename(project, fmt, draft),
        size_bytes=final_path.stat().st_size,
        sha256=sha256_file(final_path),
        draft=draft,
        errors_at_export=report.errors,
        warnings_at_export=report.unacknowledged_warnings,
        checks=checks,
        warnings=warnings,
    )


def preview_docx_bytes(store: ProjectStore, settings: Settings, project: Project) -> bytes:
    if project.template is None or project.template.format != "docx":
        raise AppError("unsupported_preview", "Template preview is available for DOCX templates.", 422)
    with tempfile.TemporaryDirectory(prefix="tf-preview-") as tmp:
        path = Path(tmp) / "preview.docx"
        generate_file(store, settings, project, "docx", True, path)
        return path.read_bytes()


def preview_xlsx_grid(store: ProjectStore, settings: Settings, project: Project) -> dict:
    if project.template is None or project.template.format != "xlsx":
        raise AppError("unsupported_preview", "Grid preview is available for XLSX templates.", 422)
    from openpyxl import load_workbook

    with tempfile.TemporaryDirectory(prefix="tf-preview-") as tmp:
        path = Path(tmp) / "preview.xlsx"
        generate_file(store, settings, project, "xlsx", True, path)
        wb = load_workbook(str(path), data_only=False)
        try:
            ws = wb.worksheets[0]
            rows = []
            for row in ws.iter_rows(
                min_row=1, max_row=min(ws.max_row, 80), max_col=min(ws.max_column, 12), values_only=True
            ):
                rows.append(["" if v is None else (f"{v:,.2f}" if isinstance(v, float) else str(v)) for v in row])
            return {"sheet": ws.title, "rows": rows}
        finally:
            wb.close()
