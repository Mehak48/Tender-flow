"""Quotation template analysis (pipeline stage 5).

Supported template conventions (documented in docs/SYSTEM.md):

* ``{{field_key}}`` placeholders anywhere in DOCX paragraphs, tables, headers and footers, or in XLSX
  cells. ``{{field_key?}}`` marks an optional field.
* A line-items table, found by its header row (e.g. "Description", "Qty", "Unit price", "Amount").
  The row directly below the header is the row that is repeated for every line item. Alternatively a
  row containing ``{{item.description}}``-style placeholders.
* "Label: ______" blanks in paragraphs, and table cells such as "Quotation No." followed by an empty
  cell. These become destination fields when the label is recognised.

Signature-type fields are detected but never filled.
"""

from __future__ import annotations

import re
import zipfile
from collections.abc import Iterator
from dataclasses import dataclass
from pathlib import Path

from docx import Document
from docx.document import Document as DocxDocument
from docx.table import Table
from docx.text.paragraph import Paragraph
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

from .errors import AppError
from .fields import (
    ITEM_ATTRIBUTES,
    QUOTATION_FIELDS,
    canonical_placeholder,
    match_field_label,
    match_item_header,
)
from .models import DocumentRecord, ItemColumn, ItemsTableInfo, TemplateAnalysis, TemplateField
from .parsing.docx_parser import unique_cells

PLACEHOLDER_RE = re.compile(r"\{\{\s*([A-Za-z][A-Za-z0-9_.]*)\s*(\?)?\s*\}\}")
ITEM_PLACEHOLDER_RE = re.compile(r"\{\{\s*item\.([a-z_]+)\s*\}\}", re.I)
BLANK_RE = re.compile(r"^(?P<label>[A-Za-z][A-Za-z0-9 .&/()#'\-]{0,40}?)\s*:\s*(?P<blank>_{3,}|\.{4,}|…+)\s*$")
TOTAL_KEYS = {"subtotal", "tax_amount", "total", "tax_rate"}


@dataclass
class _Loc:
    text: str
    location: str
    table: int | None = None  # 0-based body table index
    row: int | None = None
    col: int | None = None


class _Builder:
    def __init__(self) -> None:
        self.fields: dict[str, TemplateField] = {}
        self.warnings: list[str] = []

    def placeholder(self, raw_key: str, optional: bool, location: str) -> None:
        key_written = raw_key.strip().lower()
        if key_written.startswith("item."):
            return
        fid = f"ph:{key_written}"
        existing = self.fields.get(fid)
        if existing:
            if location not in existing.locations:
                existing.locations.append(location)
            if not optional:
                existing.required = True
            return
        canonical = canonical_placeholder(key_written)
        never_fill = bool(canonical and QUOTATION_FIELDS[canonical].never_fill)
        self.fields[fid] = TemplateField(
            id=fid,
            key=canonical,
            label=QUOTATION_FIELDS[canonical].label if canonical else raw_key,
            kind="placeholder",
            required=(not optional) and not never_fill,
            recognized=canonical is not None,
            locations=[location],
        )

    def label_blank(self, fid: str, label: str, location: str) -> None:
        key = match_field_label(label)
        if key is None or fid in self.fields:
            return
        never_fill = QUOTATION_FIELDS[key].never_fill
        self.fields[fid] = TemplateField(
            id=fid,
            key=key,
            label=label.strip().rstrip(":").strip() or QUOTATION_FIELDS[key].label,
            kind="label_blank",
            required=not never_fill,
            recognized=True,
            locations=[location],
        )


def _scan_placeholders(builder: _Builder, text: str, location: str) -> None:
    for m in PLACEHOLDER_RE.finditer(text):
        builder.placeholder(m.group(1), bool(m.group(2)), location)


def _iter_paragraph_locs(document: DocxDocument) -> Iterator[tuple[Paragraph, str]]:
    para_no = 0
    for item in document.iter_inner_content():
        if isinstance(item, Paragraph):
            para_no += 1
            yield item, f"Paragraph {para_no}"


def _items_table_from_rows(rows_text: list[list[str]]) -> tuple[int, int, list[ItemColumn], str] | None:
    # Placeholder row mode: a row containing {{item.x}} placeholders.
    for r, cells in enumerate(rows_text):
        attrs = [ITEM_PLACEHOLDER_RE.search(c) for c in cells]
        if sum(1 for a in attrs if a) >= 2:
            header_r = r - 1 if r > 0 else r
            header_cells = rows_text[header_r] if header_r != r else ["" for _ in cells]
            columns = []
            for i, a in enumerate(attrs):
                attr = a.group(1).lower() if a else None
                columns.append(
                    ItemColumn(
                        index=i,
                        header=header_cells[i] if i < len(header_cells) else "",
                        attribute=attr if attr in ITEM_ATTRIBUTES else None,
                    )
                )
            return header_r, r, columns, "placeholder_row"
    for r, cells in enumerate(rows_text[:3]):
        header_attrs: list[str | None] = [match_item_header(c) for c in cells]
        recognised = {a for a in header_attrs if a}
        if "description" in recognised and recognised & {"quantity", "unit_price", "amount"} and len(recognised) >= 2:
            seen: set[str] = set()
            columns = []
            for i, attr_name in enumerate(header_attrs):
                attribute = None if attr_name in seen else attr_name
                if attribute:
                    seen.add(attribute)
                columns.append(ItemColumn(index=i, header=cells[i], attribute=attribute))
            return r, r + 1, columns, "header_detected"
    return None


def analyze_docx(path: Path, record: DocumentRecord) -> TemplateAnalysis:
    try:
        document = Document(str(path))
    except Exception as exc:
        raise AppError("unreadable_docx", "This Word template could not be read. It may be damaged.") from exc
    b = _Builder()
    items_info: ItemsTableInfo | None = None

    for paragraph, loc in _iter_paragraph_locs(document):
        text = paragraph.text
        _scan_placeholders(b, text, loc)
        m = BLANK_RE.match(text.strip())
        if m and not PLACEHOLDER_RE.search(text):
            b.label_blank(f"lbl:{loc.replace(' ', '').lower()}", m.group("label"), loc)

    tables = [item for item in document.iter_inner_content() if isinstance(item, Table)]
    for t_index, table in enumerate(tables):
        rows = list(table.rows)
        rows_cells = [unique_cells(row) for row in rows]
        rows_text = [[c.text.strip() for c in cells] for cells in rows_cells]
        if items_info is None:
            found = _items_table_from_rows(rows_text)
            if found:
                header_r, template_r, columns, mode = found
                if template_r >= len(rows):
                    b.warnings.append(
                        "The items table has no row below its header; TenderFlow adds rows using the header row's "
                        "layout, so check the formatting of the exported table."
                    )
                items_info = ItemsTableInfo(
                    location=f"Table {t_index + 1}",
                    table_index=t_index,
                    header_row=header_r,
                    template_row=template_r,
                    columns=columns,
                    mode=mode,  # type: ignore[arg-type]
                )
        skip_rows: set[int] = set()
        if items_info is not None and items_info.table_index == t_index:
            skip_rows = {items_info.header_row, items_info.template_row}
        for r, cells in enumerate(rows_cells):
            row_has_placeholder = any(PLACEHOLDER_RE.search(c.text) for c in cells)
            for c_i, cell in enumerate(cells):
                loc = f"Table {t_index + 1}, row {r + 1}, cell {c_i + 1}"
                _scan_placeholders(b, cell.text, loc)
            if r in skip_rows or row_has_placeholder:
                continue
            texts = [c.text.strip() for c in cells]
            for c_i, text in enumerate(texts):
                if not text or c_i + 1 >= len(texts):
                    continue
                key = match_field_label(text.rstrip(":"))
                if key is None:
                    continue
                following = texts[c_i + 1 :]
                if all(not t for t in following):
                    target = len(texts) - 1 if key in TOTAL_KEYS else c_i + 1
                    b.label_blank(
                        f"lbl:t{t_index + 1}r{r + 1}c{target + 1}",
                        text,
                        f"Table {t_index + 1}, row {r + 1}, cell {target + 1}",
                    )
                    break

    seen_parts: list = []  # element references (lxml proxies are only stable while referenced)
    for section in document.sections:
        for kind, part in (
            ("Header", section.header),
            ("Footer", section.footer),
            ("First-page header", section.first_page_header),
            ("First-page footer", section.first_page_footer),
            ("Even-page header", section.even_page_header),
            ("Even-page footer", section.even_page_footer),
        ):
            if part.is_linked_to_previous:
                continue
            element = part._element
            if any(element is seen for seen in seen_parts):
                continue
            seen_parts.append(element)
            for p in part.paragraphs:
                _scan_placeholders(b, p.text, kind)
            for tbl in part.tables:
                for row in tbl.rows:
                    for cell in unique_cells(row):
                        _scan_placeholders(b, cell.text, kind)

    return _finish(b, record, "docx", items_info)


def analyze_xlsx(path: Path, record: DocumentRecord) -> TemplateAnalysis:
    try:
        wb = load_workbook(str(path), read_only=False, data_only=False)
    except (zipfile.BadZipFile, KeyError, ValueError, OSError) as exc:
        raise AppError("unreadable_xlsx", "This Excel template could not be read. It may be damaged.") from exc
    b = _Builder()
    items_info: ItemsTableInfo | None = None
    try:
        for ws in wb.worksheets:
            max_row = min(ws.max_row, 500)
            max_col = min(ws.max_column, 40)
            grid: dict[tuple[int, int], str] = {}
            for row in ws.iter_rows(min_row=1, max_row=max_row, max_col=max_col):
                for cell in row:
                    if isinstance(cell.value, str) and cell.value.strip():
                        grid[(cell.row, cell.column)] = cell.value.strip()
            for (r, c), text in sorted(grid.items()):
                _scan_placeholders(b, text, f"{ws.title}!{get_column_letter(c)}{r}")
            if items_info is None:
                for r in range(1, max_row + 1):
                    cells = [grid.get((r, c), "") for c in range(1, max_col + 1)]
                    attrs = [match_item_header(t) if t else None for t in cells]
                    recognised = {a for a in attrs if a}
                    if "description" in recognised and recognised & {"quantity", "unit_price", "amount"}:
                        last = max(i for i, a in enumerate(attrs) if a is not None)
                        columns: list[ItemColumn] = []
                        seen: set[str] = set()
                        for i in range(0, last + 1):
                            a = attrs[i] if attrs[i] not in seen else None
                            if a:
                                seen.add(a)
                            columns.append(ItemColumn(index=i, header=cells[i], attribute=a))
                        items_info = ItemsTableInfo(
                            location=f"Sheet {ws.title}, rows {r + 1}+",
                            sheet=ws.title,
                            header_row=r,
                            template_row=r + 1,
                            columns=columns,
                            mode="header_detected",
                        )
                        break
            header_row = items_info.header_row if items_info and items_info.sheet == ws.title else None
            for (r, c), text in sorted(grid.items()):
                if r == header_row or PLACEHOLDER_RE.search(text):
                    continue
                row_has_placeholder = any(PLACEHOLDER_RE.search(grid.get((r, cc), "")) for cc in range(1, max_col + 1))
                if row_has_placeholder:
                    continue
                key = match_field_label(text.rstrip(":"))
                if key is None:
                    continue
                if not grid.get((r, c + 1)) and c + 1 <= max_col + 1:
                    ref = f"{get_column_letter(c + 1)}{r}"
                    b.label_blank(f"lbl:{ws.title}!{ref}", text, f"{ws.title}!{ref}")
    finally:
        wb.close()
    return _finish(b, record, "xlsx", items_info)


def _finish(b: _Builder, record: DocumentRecord, fmt: str, items_info: ItemsTableInfo | None) -> TemplateAnalysis:
    warnings = list(b.warnings)
    if items_info is None:
        warnings.append(
            "No line-items table was found. Add a table with headers such as “Description”, “Qty”, "
            "“Unit price” and “Amount” so line items can be written into the template."
        )
    elif not any(c.attribute == "description" for c in items_info.columns):
        warnings.append("The items table has no recognisable description column.")
    if not b.fields:
        warnings.append(
            "No placeholders or fill-in labels were found. Add placeholders like {{company_name}} or "
            "“Quotation No.: ______” so values can be filled in."
        )
    unknown = [f.label for f in b.fields.values() if not f.recognized]
    if unknown:
        warnings.append("Unrecognised placeholders (map them manually or leave them empty): " + ", ".join(unknown[:8]))
    if any(f.key == "signature" for f in b.fields.values()):
        warnings.append(
            "A signature field was found. TenderFlow never fills signatures; sign the final document yourself."
        )
    return TemplateAnalysis(
        document_id=record.id,
        document_name=record.filename,
        format=fmt,  # type: ignore[arg-type]
        fields=list(b.fields.values()),
        items_table=items_info,
        warnings=warnings,
    )


def analyze_template(path: Path, record: DocumentRecord) -> TemplateAnalysis:
    if record.format == "docx":
        return analyze_docx(path, record)
    if record.format == "xlsx":
        return analyze_xlsx(path, record)
    raise AppError("unsupported_format", "Templates must be .docx or .xlsx files.", 415)
