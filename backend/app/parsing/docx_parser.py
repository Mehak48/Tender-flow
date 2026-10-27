"""DOCX parsing with python-docx: body paragraphs and tables in document order.

Paragraph numbers count body paragraphs (1-based) the way a reader would scan the document; tables are
numbered document-wide (1-based) and rows within a table are 1-based. Horizontally merged cells are
de-duplicated so a merged cell is read once.
"""

from __future__ import annotations

import zipfile
from pathlib import Path

from docx import Document
from docx.opc.exceptions import PackageNotFoundError
from docx.table import Table
from docx.text.paragraph import Paragraph

from ..errors import AppError
from .types import Block, ParsedDocument, ParsedTable, clean_text

MAX_BLOCKS = 20000


def unique_cells(row) -> list:
    """Return the cells of a table row with horizontally merged cells collapsed into one."""
    seen: list = []
    for cell in row.cells:
        if seen and cell._tc is seen[-1]._tc:
            continue
        seen.append(cell)
    return seen


def parse_docx(path: Path) -> ParsedDocument:
    try:
        document = Document(str(path))
    except (PackageNotFoundError, zipfile.BadZipFile, KeyError, ValueError) as exc:
        raise AppError("unreadable_docx", "This Word document could not be read. It may be damaged.") from exc
    except Exception as exc:  # lxml and python-docx can raise several types for malformed XML
        raise AppError("unreadable_docx", "This Word document could not be read. It may be damaged.") from exc

    doc = ParsedDocument(format="docx", page_count=None)
    paragraph_no = 0
    table_no = 0
    for item in document.iter_inner_content():
        if len(doc.blocks) > MAX_BLOCKS:
            doc.warnings.append("The document is very long; only the first part was read.")
            break
        if isinstance(item, Paragraph):
            text = clean_text(item.text)
            if not text:
                continue
            paragraph_no += 1
            style = (item.style.name if item.style is not None else "") or ""
            kind = "heading" if style.lower().startswith(("heading", "title")) else "paragraph"
            doc.blocks.append(Block(id=f"para-{paragraph_no}", kind=kind, text=text, paragraph=paragraph_no))
        elif isinstance(item, Table):
            table_no += 1
            rows: list[list[str]] = []
            block_ids: list[str] = []
            for r_index, row in enumerate(item.rows, start=1):
                cells = [clean_text(c.text) for c in unique_cells(row)]
                if not any(cells):
                    continue
                block_id = f"t{table_no}-r{r_index}"
                block_ids.append(block_id)
                rows.append(cells)
                doc.blocks.append(
                    Block(
                        id=block_id,
                        kind="table_row",
                        text=" | ".join(c for c in cells if c),
                        table=table_no,
                        row=r_index,
                        cells=cells,
                    )
                )
            if rows:
                doc.tables.append(ParsedTable(index=table_no, rows=rows, block_ids=block_ids))
    if not doc.blocks:
        doc.warnings.append("No text was found in this Word document.")
    return doc
